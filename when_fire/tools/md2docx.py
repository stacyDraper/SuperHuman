"""Convert the design-intent markdown (a small, known subset) to a real .docx with python-docx."""
import re
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, OUT = sys.argv[1], sys.argv[2]

doc = Document()

# Page: US Letter, 1 inch margins
for s in doc.sections:
    s.page_width = Inches(8.5)
    s.page_height = Inches(11)
    s.left_margin = s.right_margin = Inches(1)
    s.top_margin = s.bottom_margin = Inches(1)

# Base styles
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

for name, size, color in (("Title", 24, "1F3864"), ("Heading 1", 16, "1F3864"), ("Heading 2", 13, "2E5597")):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if name != "Title":
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True

lb = doc.styles["List Bullet"]
lb.font.name = "Calibri"
lb.font.size = Pt(11)
lb.paragraph_format.space_after = Pt(3)


def add_runs(par, text):
    """Inline **bold** handling.  Odd segments between ** markers are bold."""
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if not seg:
            continue
        r = par.add_run(seg)
        if i % 2 == 1:
            r.bold = True


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(inches * 1440)))
    tcW.set(qn("w:type"), "dxa")


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def add_table(rows):
    ncols = len(rows[0])
    widths = {2: [1.45, 5.05], 3: [3.0, 2.25, 1.25]}.get(ncols, [6.5 / ncols] * ncols)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for ci, col in enumerate(t.columns):
        col.width = Inches(widths[ci])
    header_row = ncols == 3  # the ledger has a header row; the metadata table does not
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, txt in enumerate(row):
            set_cell_width(cells[ci], widths[ci])
            p = cells[ci].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, txt.strip())
            for r in p.runs:
                r.font.size = Pt(10)
            if (header_row and ri == 0) or (ncols == 2):
                if ci == 0 or (header_row and ri == 0):
                    shade(cells[ci], "E8EEF7")
    doc.add_paragraph()  # breathing room after a table


def flush_table(buf):
    rows = []
    for line in buf:
        inner = line.strip().strip("|")
        cells = [c for c in inner.split("|")]
        if all(re.fullmatch(r"\s*-{3,}\s*", c) for c in cells):
            continue
        rows.append(cells)
    if rows:
        add_table(rows)


with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

table_buf = []
for raw in lines:
    line = raw.rstrip()
    if line.startswith("|"):
        table_buf.append(line)
        continue
    if table_buf:
        flush_table(table_buf)
        table_buf = []
    if not line.strip():
        continue
    if line.startswith("# "):
        p = doc.add_paragraph(style="Title")
        add_runs(p, line[2:])
    elif line.startswith("## "):
        p = doc.add_heading(level=1)
        add_runs(p, line[3:])
    elif line.startswith("### "):
        p = doc.add_heading(level=2)
        add_runs(p, line[4:])
    elif line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, line[2:])
    else:
        p = doc.add_paragraph()
        add_runs(p, line)
if table_buf:
    flush_table(table_buf)

# Header and footer with page number
sec = doc.sections[0]
hp = sec.header.paragraphs[0]
title_line = next((l[2:].strip() for l in lines if l.startswith("# ")), "")
hp.text = sys.argv[3] if len(sys.argv) > 3 else f"{title_line}  ·  living document"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs:
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string("666666")

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tag)
        run._r.append(el)
    else:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = text
        run._r.append(el)
run.font.size = Pt(9)

# python-docx's template ships a <w:zoom> without the required percent attribute
zoom = doc.settings.element.find(qn("w:zoom"))
if zoom is not None and zoom.get(qn("w:percent")) is None:
    zoom.set(qn("w:percent"), "100")

doc.save(OUT)

# Slim: drop the optional effects-styles part and the thumbnail the template ships with.
import zipfile, os
tmp = OUT + ".slim"
drop = {"word/stylesWithEffects.xml", "docProps/thumbnail.jpeg", "word/_rels/stylesWithEffects.xml.rels"}
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        if item.filename in drop:
            continue
        data = zin.read(item.filename)
        if item.filename == "[Content_Types].xml":
            s = data.decode("utf-8")
            s = re.sub(r'<Override[^>]*stylesWithEffects[^>]*/>', '', s)
            s = re.sub(r'<Default Extension="jpeg"[^>]*/>', '', s)
            data = s.encode("utf-8")
        elif item.filename == "word/_rels/document.xml.rels":
            s = data.decode("utf-8")
            s = re.sub(r'<Relationship[^>]*stylesWithEffects[^>]*/>', '', s)
            data = s.encode("utf-8")
        elif item.filename == "_rels/.rels":
            s = data.decode("utf-8")
            s = re.sub(r'<Relationship[^>]*thumbnail[^>]*/>', '', s)
            data = s.encode("utf-8")
        zout.writestr(item, data)
os.replace(tmp, OUT)
print("saved", OUT)
